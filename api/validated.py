"""Excel export for the complete candidate review captured by the search UI."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


HEADERS = (
    "Código MAI",
    "Prestación MAI",
    "ConceptID SNOMED CT",
    "Término preferido SNOMED CT (es)",
    "Término coincidente",
    "Etiqueta semántica",
    "Estado de revisión",
    "Motivo de no selección",
    "Detalle de otra justificación",
    "Referencia bibliográfica",
    "Motivo de selección",
    "Observaciones adicionales",
    "Notas",
    "Fecha de revisión",
)

V2_HEADERS = (
    "Versión del mapeo",
    "Código MAI",
    "Display MAI",
    "Grupo de componentes",
    "Fuente OCL",
    "Código de examen OCL",
    "Examen único OCL (es)",
    "Examen único OCL (en)",
    "Componente obligatorio",
    "Mapeo del componente confirmado",
    "ConceptID SNOMED CT",
    "Término preferido SNOMED CT (es)",
    "Estado de revisión",
    "Motivo de no selección",
    "Detalle de otra justificación",
    "Referencia bibliográfica",
    "Motivo de selección",
    "Observaciones adicionales",
    "Fecha de revisión",
)


def _excel_text(value: object) -> str:
    """Keep identifiers as text and neutralize formula-like user input."""
    text = "" if value is None else str(value)
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def create_validated_mappings_xlsx(mappings: list[dict]) -> bytes:
    """Return a formatted one-row-per-reviewed-candidate workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Revisión de mapeos"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.append(HEADERS)

    for item in mappings:
        ws.append([
            _excel_text(item.get("source_code")),
            _excel_text(item.get("source_term")),
            _excel_text(item.get("concept_id")),
            _excel_text(item.get("fsn")),
            _excel_text(item.get("matched_term")),
            _excel_text(item.get("semantic_tag")),
            _excel_text(item.get("status")),
            _excel_text(item.get("rejection_reason")),
            _excel_text(item.get("rejection_detail")),
            _excel_text(item.get("reference")),
            _excel_text(item.get("selection_reason")),
            _excel_text(item.get("observations")),
            _excel_text(item.get("notes")),
            _excel_text(item.get("reviewed_at")),
        ])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.number_format = "@"

    widths = (16, 38, 24, 48, 38, 22, 20, 34, 38, 45, 45, 45, 36, 24)
    for column, width in zip(ws.columns, widths):
        ws.column_dimensions[column[0].column_letter].width = width
    ws.row_dimensions[1].height = 30

    last_row = max(ws.max_row, 2)
    if ws.max_row == 1:  # Tables need a data row; the endpoint rejects this case, kept for safety.
        ws.append([""] * len(HEADERS))
    table = Table(displayName="RevisionMapeos", ref=f"A1:N{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def create_mapping_v2_xlsx(mappings: list[dict]) -> bytes:
    """Export one auditable row per OCL component/SNOMED candidate in mapping v2."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Mapeo MAI-SNOMED v2"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.append(V2_HEADERS)
    for item in mappings:
        ws.append([
            _excel_text(item.get("mapping_version") or "2.0.1"),
            _excel_text(item.get("mai_code") or item.get("source_code")),
            _excel_text(item.get("mai_display")),
            _excel_text(item.get("component_group") or item.get("mai_code")),
            _excel_text(item.get("ocl_source") or "MINSALQA/manualimagenes"),
            _excel_text(item.get("ocl_concept_id")),
            _excel_text(item.get("ocl_display") or item.get("source_term")),
            _excel_text(item.get("ocl_display_en")),
            "Sí" if item.get("required_component", True) else "No",
            "Sí" if item.get("mapping_confirmed", False) else "No",
            _excel_text(item.get("concept_id")),
            _excel_text(item.get("fsn")),
            _excel_text(item.get("status")),
            _excel_text(item.get("rejection_reason")),
            _excel_text(item.get("rejection_detail")),
            _excel_text(item.get("reference")),
            _excel_text(item.get("selection_reason")),
            _excel_text(item.get("observations")),
            _excel_text(item.get("reviewed_at")),
        ])

    header_fill = PatternFill("solid", fgColor="385723")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.number_format = "@"
    widths = (16, 16, 42, 20, 26, 22, 48, 46, 20, 24, 24, 48, 20, 34, 38, 44, 44, 44, 24)
    for column, width in zip(ws.columns, widths):
        ws.column_dimensions[column[0].column_letter].width = width
    ws.row_dimensions[1].height = 42
    if ws.max_row == 1:
        ws.append([""] * len(V2_HEADERS))
    table = Table(displayName="MapeoMAISNOMEDV2", ref=f"A1:S{max(ws.max_row, 2)}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
