from io import BytesIO

from openpyxl import load_workbook

from api.ocl import _component, _workspace_key
from api.validated import V2_HEADERS, create_mapping_v2_xlsx


def test_ocl_component_prefers_spanish_and_keeps_english():
    item = {
        "id": "3898",
        "display_name": "Fallback",
        "concept_class": "Radiology/Imaging Procedure",
        "names": [
            {"name": "Full field digital mammography", "locale": "en", "locale_preferred": True},
            {"name": "Mamografía digital de campo completo", "locale": "es", "locale_preferred": True},
        ],
    }
    result = _component(item)
    assert result["id"] == "3898"
    assert result["display"] == "Mamografía digital de campo completo"
    assert result["display_en"] == "Full field digital mammography"


def test_workspace_key_normalizes_case_and_spaces():
    assert _workspace_key("0401010", "Mamografía   bilateral") == _workspace_key(
        "0401010", "  MAMOGRAFÍA bilateral  "
    )


def test_v2_export_contains_grouping_and_component_columns():
    content = create_mapping_v2_xlsx([
        {
            "mapping_version": "2.0.2",
            "mai_code": "0401010",
            "mai_display": "Mamografía bilateral",
            "component_group": "0401010",
            "ocl_source": "MINSALQA/manualimagenes",
            "ocl_concept_id": "3898",
            "ocl_display": "Mamografía digital bilateral diagnóstica",
            "ocl_display_en": "Bilateral diagnostic digital mammography",
            "required_component": True,
            "mapping_confirmed": True,
            "concept_id": "43204002",
            "fsn": "mamografía de mamas bilateral",
            "status": "Seleccionado",
            "selection_reason": "Coincide con el examen único",
        }
    ])
    ws = load_workbook(BytesIO(content), data_only=False).active
    assert tuple(cell.value for cell in ws[1]) == V2_HEADERS
    assert ws["A2"].value == "2.0.2"
    assert ws["B2"].value == "0401010"
    assert ws["F2"].value == "3898"
    assert ws["J2"].value == "Sí"
    assert ws["K2"].value == "43204002"
    assert ws["Q2"].value == "Coincide con el examen único"
    assert "MapeoMAISNOMEDV2" in ws.tables
