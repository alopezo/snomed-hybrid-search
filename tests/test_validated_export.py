from io import BytesIO

from openpyxl import load_workbook

from api.validated import HEADERS, create_validated_mappings_xlsx


def test_validated_mappings_xlsx_is_auditable_long_table():
    content = create_validated_mappings_xlsx([
        {
            "source_code": "MAI-001",
            "source_term": "Radiografía de las glándulas salivales (sialografía)",
            "concept_id": "1343566005",
            "fsn": "Fluoroscopy of salivary apparatus with contrast (procedure)",
            "matched_term": "Fluoroscopic sialography with contrast",
            "semantic_tag": "procedure",
            "status": "Seleccionado",
            "reference": "Guía clínica 2025, p. 18",
            "selection_reason": "Coincide con la técnica descrita.",
            "context": "Sialografía convencional; excluir CT, RM y US.",
            "notes": "Validación manual",
            "reviewed_at": "2026-10-02T12:00:00Z",
        },
        {
            "source_code": "MAI-001",
            "source_term": "Radiografía de las glándulas salivales (sialografía)",
            "concept_id": "420088004",
            "fsn": "Fluoroscopy of salivary gland (procedure)",
            "status": "Descartado",
            "rejection_reason": "Concepto muy genérico",
        },
    ])

    wb = load_workbook(BytesIO(content), data_only=False)
    ws = wb["Revisión de mapeos"]
    assert tuple(c.value for c in ws[1]) == HEADERS
    assert ws.max_row == 3
    assert ws["A2"].value == "MAI-001"
    assert ws["C2"].value == "1343566005"
    assert ws["G2"].value == "Seleccionado"
    assert ws["J2"].value == "Guía clínica 2025, p. 18"
    assert ws["K2"].value == "Coincide con la técnica descrita."
    assert "Contexto de búsqueda" not in HEADERS
    assert ws["G3"].value == "Descartado"
    assert ws["H3"].value == "Concepto muy genérico"
    assert ws.freeze_panes == "A2"
    assert "RevisionMapeos" in ws.tables


def test_validated_mappings_xlsx_neutralizes_formula_input():
    content = create_validated_mappings_xlsx([{
        "source_term": "=HYPERLINK(\"https://example.test\")",
        "concept_id": "123456",
        "fsn": "Safe concept (procedure)",
    }])
    ws = load_workbook(BytesIO(content), data_only=False).active
    assert ws["B2"].data_type == "s"
    assert ws["B2"].value.startswith("'=")
